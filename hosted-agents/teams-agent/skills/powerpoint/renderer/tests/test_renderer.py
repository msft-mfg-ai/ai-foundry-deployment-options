import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from pptx import Presentation
from lxml import etree

from zava_renderer import DeckSpecError, render_deck


ROOT = Path(__file__).parents[1]
TEMPLATE = ROOT.parent / "assets" / "template.pptx"


def sample_spec():
    return {
        "title": "Smart Textile Safety",
        "audience": "Product and safety leaders",
        "objective": "Align on design controls",
        "slides": [
            {
                "id": "01",
                "role": "cover",
                "composition": "cover",
                "title": "Smart Textile Safety",
                "subtitle": "Controls from design through end of life",
                "body": [],
            },
            {
                "id": "02",
                "role": "summary",
                "composition": "card_grid",
                "title": "Four controls reduce risk",
                "body": [
                    {
                        "kind": "bullet",
                        "title": "Materials",
                        "text": "Verify skin compatibility, wash durability, and abrasion resistance.",
                    },
                    {
                        "kind": "bullet",
                        "title": "Power",
                        "text": "Limit current and temperature; isolate batteries and conductors.",
                    },
                    {
                        "kind": "bullet",
                        "title": "Data",
                        "text": "Use secure pairing, encryption, and minimal telemetry collection.",
                    },
                    {
                        "kind": "bullet",
                        "title": "Lifecycle",
                        "text": "Provide cleaning, repair, traceability, and disposal instructions.",
                    },
                ],
            },
        ],
    }


def test_renders_clean_zava_deck(tmp_path):
    destination = tmp_path / "deck.pptx"

    render_deck(sample_spec(), TEMPLATE, destination)

    presentation = Presentation(destination)
    assert len(presentation.slides) == 2
    assert presentation.slides[0].slide_layout.name == "Title 1"
    assert presentation.slides[1].slide_layout.name == "Title Only (Original)"
    assert presentation.slides[0].shapes.title.text == "Smart Textile Safety"
    assert presentation.slides[1].shapes.title.text == "Four controls reduce risk"
    assert all(
        not shape.has_text_frame or shape.text.strip()
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.is_placeholder
    )


def test_card_grid_stays_below_title_region(tmp_path):
    destination = tmp_path / "deck.pptx"
    render_deck(sample_spec(), TEMPLATE, destination)
    slide = Presentation(destination).slides[1]
    title = slide.shapes.title
    content_shapes = [shape for shape in slide.shapes if not shape.is_placeholder]

    assert content_shapes
    assert min(shape.top for shape in content_shapes) >= title.top + title.height
    assert max(shape.top + shape.height for shape in content_shapes) <= 5.5 * 914400


def test_package_has_one_unique_relationship_per_slide(tmp_path):
    destination = tmp_path / "deck.pptx"
    render_deck(sample_spec(), TEMPLATE, destination)

    with ZipFile(destination) as archive:
        presentation = etree.fromstring(archive.read("ppt/presentation.xml"))
        relationships = etree.fromstring(
            archive.read("ppt/_rels/presentation.xml.rels")
        )

    namespaces = {
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    active_ids = presentation.xpath("//p:sldId/@r:id", namespaces=namespaces)
    slide_relationships = [
        relationship
        for relationship in relationships.xpath(
            "//pr:Relationship",
            namespaces=namespaces,
        )
        if relationship.get("Type", "").endswith("/slide")
    ]
    relationships_by_id = {
        relationship.get("Id"): relationship.get("Target")
        for relationship in slide_relationships
    }
    active_targets = [relationships_by_id.get(active_id) for active_id in active_ids]

    assert len(active_ids) == 2
    assert len(slide_relationships) == len(active_ids)
    assert None not in active_targets
    assert len(set(active_targets)) == len(active_targets)
    assert set(active_ids) == set(relationships_by_id)


def test_rejects_oversized_card_grid(tmp_path):
    spec = sample_spec()
    spec["slides"][1]["body"].append(
        {"kind": "bullet", "title": "Extra", "text": "Too many cards"}
    )

    with pytest.raises(DeckSpecError, match="two to four cards"):
        render_deck(spec, TEMPLATE, tmp_path / "deck.pptx")


def test_rejects_non_string_subtitle(tmp_path):
    spec = sample_spec()
    spec["slides"][0]["subtitle"] = None

    with pytest.raises(DeckSpecError, match="subtitle must be a string"):
        render_deck(spec, TEMPLATE, tmp_path / "deck.pptx")


def test_example_spec_is_valid():
    example = ROOT / "examples" / "smart-textile-safety.json"
    with example.open(encoding="utf-8") as stream:
        spec = json.load(stream)

    assert spec == sample_spec()
