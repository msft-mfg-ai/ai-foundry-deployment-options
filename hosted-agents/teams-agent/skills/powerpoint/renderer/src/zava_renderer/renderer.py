from __future__ import annotations

from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide
from pptx.util import Inches, Pt


INK = RGBColor(0x0A, 0x0C, 0x0C)
DEEP_TEAL = RGBColor(0x18, 0x3D, 0x4C)
TEAL = RGBColor(0x3D, 0x72, 0x88)
SKY = RGBColor(0x97, 0xCC, 0xE3)
PALE_BLUE = RGBColor(0xBB, 0xD5, 0xE1)
LIGHT_GRAY = RGBColor(0xEB, 0xEF, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MID_TEAL = RGBColor(0x2F, 0x62, 0x78)
DARK_SKY = RGBColor(0x37, 0x67, 0x79)

SYSTEM_PLACEHOLDERS = {
    PP_PLACEHOLDER.DATE,
    PP_PLACEHOLDER.FOOTER,
    PP_PLACEHOLDER.SLIDE_NUMBER,
}


class DeckSpecError(ValueError):
    pass


def render_deck(
    spec: Mapping[str, Any],
    template_path: Path | str,
    output_path: Path | str,
) -> Path:
    normalized = _validate_spec(spec)
    presentation = Presentation(str(template_path))
    if presentation.slides:
        raise DeckSpecError("The renderer requires the clean zero-slide Zava template.")

    for slide_spec in normalized["slides"]:
        composition = slide_spec["composition"]
        if composition == "cover":
            _render_cover(presentation, slide_spec)
        elif composition == "card_grid":
            _render_card_grid(presentation, slide_spec)
        else:
            raise DeckSpecError(f"Unsupported composition: {composition}")

    _validate_presentation(presentation)
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(destination)

    reopened = Presentation(str(destination))
    _validate_presentation(reopened)
    return destination


def _validate_spec(spec: Mapping[str, Any]) -> dict[str, Any]:
    if not isinstance(spec, Mapping):
        raise DeckSpecError("DeckSpec must be a JSON object.")

    slides = spec.get("slides")
    if not isinstance(slides, Sequence) or isinstance(slides, (str, bytes)):
        raise DeckSpecError("DeckSpec.slides must be an array.")
    if not slides:
        raise DeckSpecError("DeckSpec must contain at least one slide.")

    normalized_slides: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    for index, raw_slide in enumerate(slides, start=1):
        if not isinstance(raw_slide, Mapping):
            raise DeckSpecError(f"Slide {index} must be an object.")
        slide_id = _required_text(raw_slide, "id", f"Slide {index}")
        if slide_id in seen_ids:
            raise DeckSpecError(f"Duplicate slide id: {slide_id}")
        seen_ids.add(slide_id)

        composition = _required_text(raw_slide, "composition", f"Slide {slide_id}")
        title = _required_text(raw_slide, "title", f"Slide {slide_id}")
        body = raw_slide.get("body", [])
        if not isinstance(body, Sequence) or isinstance(body, (str, bytes)):
            raise DeckSpecError(f"Slide {slide_id}.body must be an array.")

        if composition == "cover":
            if len(title) > 60:
                raise DeckSpecError("Cover titles must be 60 characters or fewer.")
            raw_subtitle = raw_slide.get("subtitle", "")
            if not isinstance(raw_subtitle, str):
                raise DeckSpecError(f"Slide {slide_id}.subtitle must be a string.")
            subtitle = raw_subtitle.strip()
            normalized_slides.append(
                {
                    "id": slide_id,
                    "composition": composition,
                    "title": title,
                    "subtitle": subtitle,
                }
            )
            continue

        if composition == "card_grid":
            if len(title.split()) > 8:
                raise DeckSpecError("Card-grid titles must be eight words or fewer.")
            if not 2 <= len(body) <= 4:
                raise DeckSpecError("Card grids require two to four cards.")
            cards = [_validate_card(item, slide_id) for item in body]
            normalized_slides.append(
                {
                    "id": slide_id,
                    "composition": composition,
                    "title": title,
                    "body": cards,
                }
            )
            continue

        raise DeckSpecError(f"Unsupported composition: {composition}")

    return {"slides": normalized_slides}


def _validate_card(item: Any, slide_id: str) -> dict[str, str]:
    if not isinstance(item, Mapping) or item.get("kind") != "bullet":
        raise DeckSpecError(
            f"Slide {slide_id} card-grid entries must be bullet objects."
        )
    title = _required_text(item, "title", f"Slide {slide_id} card")
    text = _required_text(item, "text", f"Slide {slide_id} card")
    if len(title) > 32:
        raise DeckSpecError("Card titles must be 32 characters or fewer.")
    if len(text) > 140:
        raise DeckSpecError("Card text must be 140 characters or fewer.")
    return {"title": title, "text": text}


def _required_text(source: Mapping[str, Any], key: str, scope: str) -> str:
    value = source.get(key)
    if not isinstance(value, str) or not value.strip():
        raise DeckSpecError(f"{scope}.{key} is required.")
    return value.strip()


def _layout(presentation: PresentationType, name: str):
    try:
        return next(layout for layout in presentation.slide_layouts if layout.name == name)
    except StopIteration as exc:
        raise DeckSpecError(f"Template layout not found: {name}") from exc


def _render_cover(
    presentation: PresentationType,
    slide_spec: Mapping[str, Any],
) -> None:
    slide = presentation.slides.add_slide(_layout(presentation, "Title 1"))
    title = slide.shapes.title
    if title is None:
        raise DeckSpecError("Title 1 does not contain a title placeholder.")
    title.text = slide_spec["title"]
    _format_text_frame(title.text_frame, 36, WHITE, bold=True)

    subtitle = str(slide_spec.get("subtitle", "")).strip()
    subtitle_shape = next(
        (
            shape
            for shape in slide.placeholders
            if shape.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE
        ),
        None,
    )
    if subtitle_shape is not None:
        if subtitle:
            subtitle_shape.text = subtitle
            _format_text_frame(subtitle_shape.text_frame, 18, WHITE)
        else:
            _remove_shape(subtitle_shape)

    _remove_unused_placeholders(slide)


def _render_card_grid(
    presentation: PresentationType,
    slide_spec: Mapping[str, Any],
) -> None:
    slide = presentation.slides.add_slide(
        _layout(presentation, "Title Only (Original)")
    )
    title = slide.shapes.title
    if title is None:
        raise DeckSpecError("Title Only (Original) has no title placeholder.")
    title.text = slide_spec["title"]
    _format_text_frame(title.text_frame, 32, DEEP_TEAL, bold=True)

    cards = slide_spec["body"]
    count = len(cards)
    left = 0.55
    right = 0.55
    gap = 0.28
    top = 2.05
    height = 3.45
    width = (13.333 - left - right - (gap * (count - 1))) / count
    accents = (TEAL, MID_TEAL, DEEP_TEAL, DARK_SKY)

    for index, card in enumerate(cards):
        x = left + index * (width + gap)
        accent = accents[index]
        _add_card(slide, x, top, width, height, accent, card)

    _remove_unused_placeholders(slide)


def _add_card(
    slide: Slide,
    x: float,
    y: float,
    width: float,
    height: float,
    accent: RGBColor,
    card: Mapping[str, str],
) -> None:
    panel = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = LIGHT_GRAY
    panel.line.color.rgb = accent
    panel.line.width = Pt(1.5)
    panel.adjustments[0] = 0.08

    header = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(0.72),
    )
    header.fill.solid()
    header.fill.fore_color.rgb = accent
    header.line.color.rgb = accent
    header.adjustments[0] = 0.08

    header.text = card["title"]
    header.text_frame.margin_left = Inches(0.18)
    header.text_frame.margin_right = Inches(0.12)
    header.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _format_text_frame(header.text_frame, 18, WHITE, bold=True)

    body = slide.shapes.add_textbox(
        Inches(x + 0.18),
        Inches(y + 0.95),
        Inches(width - 0.36),
        Inches(height - 1.12),
    )
    body.text = card["text"]
    body.text_frame.word_wrap = True
    body.text_frame.margin_left = 0
    body.text_frame.margin_right = 0
    body.text_frame.margin_top = 0
    body.text_frame.margin_bottom = 0
    _format_text_frame(body.text_frame, 16, INK)


def _format_text_frame(
    text_frame,
    size: int,
    color: RGBColor,
    *,
    bold: bool = False,
) -> None:
    text_frame.word_wrap = True
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color


def _remove_unused_placeholders(slide: Slide) -> None:
    for shape in list(slide.placeholders):
        if shape.placeholder_format.type in SYSTEM_PLACEHOLDERS:
            continue
        if shape.has_text_frame and shape.text.strip():
            continue
        _remove_shape(shape)


def _remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def _validate_presentation(presentation: PresentationType) -> None:
    if not presentation.slides:
        raise DeckSpecError("The rendered presentation has no slides.")

    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0:
                raise DeckSpecError(f"Slide {slide_number} has a shape outside its bounds.")
            if shape.left + shape.width > presentation.slide_width:
                raise DeckSpecError(f"Slide {slide_number} has a shape beyond its width.")
            if shape.top + shape.height > presentation.slide_height:
                raise DeckSpecError(f"Slide {slide_number} has a shape beyond its height.")
            if not shape.is_placeholder:
                continue
            if shape.placeholder_format.type in SYSTEM_PLACEHOLDERS:
                continue
            if not shape.has_text_frame or not shape.text.strip():
                raise DeckSpecError(
                    f"Slide {slide_number} contains an unused content placeholder."
                )
